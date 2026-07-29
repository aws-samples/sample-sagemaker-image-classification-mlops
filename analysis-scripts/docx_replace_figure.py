#!/usr/bin/env python3
# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0

import argparse
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu
from PIL import Image

# Word EMU (English Metric Units): 914,400 per inch.
EMU_PER_INCH = 914_400


def list_figures(docx_path: Path) -> None:
    """Print a map of inline images with nearby caption text."""
    doc = Document(str(docx_path))
    images_with_context = []
    for i, para in enumerate(doc.paragraphs):
        has_image = any(
            "graphic" in run._element.xml.lower() or "pic:pic" in run._element.xml.lower()
            for run in para.runs
        )
        if has_image:
            next_text = doc.paragraphs[i + 1].text.strip() if i + 1 < len(doc.paragraphs) else ""
            images_with_context.append(next_text[:120])

    print(f"Total inline images: {len(doc.inline_shapes)}")
    print()
    print("Figure N (1-based) | caption text")
    print("-" * 80)
    for idx, caption in enumerate(images_with_context, start=1):
        print(f"  {idx:<3} | {caption}")


def _page_content_width_emu(doc) -> int:
    """Return the usable content width (page width - left/right margins) in EMU."""
    section = doc.sections[0]
    return section.page_width - section.left_margin - section.right_margin


def replace_figure(
    docx_path: Path,
    figure_index: int,
    new_image_path: Path,
    backup: bool = True,
    width_inches: float | None = None,
    keep_size: bool = False,
) -> None:
    if figure_index < 1:
        raise ValueError("figure-index is 1-based; use 1, 2, 3, ...")

    if backup:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = docx_path.with_suffix(docx_path.suffix + f".backup_{ts}")
        shutil.copy(docx_path, backup_path)
        print(f"  backup: {backup_path}")

    doc = Document(str(docx_path))
    shapes = doc.inline_shapes
    if figure_index > len(shapes):
        raise IndexError(
            f"Document has {len(shapes)} inline images; figure-index {figure_index} is out of range"
        )

    shape = shapes[figure_index - 1]
    blip = shape._inline.graphic.graphicData.pic.blipFill.blip
    r_id = blip.get(qn("r:embed"))
    rel = doc.part.related_parts[r_id]

    new_bytes = new_image_path.read_bytes()
    rel._blob = new_bytes

    # Read native image dimensions
    with Image.open(new_image_path) as im:
        img_w_px, img_h_px = im.size
    aspect_ratio = img_w_px / img_h_px

    old_w_emu = shape.width
    old_h_emu = shape.height

    if keep_size:
        new_w_emu, new_h_emu = old_w_emu, old_h_emu
    else:
        # Default: fit to page content width
        if width_inches is not None:
            target_w_emu = int(width_inches * EMU_PER_INCH)
        else:
            target_w_emu = _page_content_width_emu(doc)
        new_w_emu = target_w_emu
        new_h_emu = int(target_w_emu / aspect_ratio)
        shape.width = Emu(new_w_emu)
        shape.height = Emu(new_h_emu)

    doc.save(str(docx_path))

    print(f"  Figure {figure_index}: {rel.partname}")
    print(f"  Source:        {new_image_path}")
    print(f"  Native image:  {img_w_px}x{img_h_px} px  (aspect {aspect_ratio:.2f})")
    print(f"  Bytes:         {len(new_bytes):,}")
    print(f"  Old in-doc:    {old_w_emu // 9525}x{old_h_emu // 9525} px")
    print(
        f"  New in-doc:    {new_w_emu // 9525}x{new_h_emu // 9525} px  ({new_w_emu / EMU_PER_INCH:.2f} in wide)"
    )
    print(f"  ✓ Saved {docx_path}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Replace an inline figure in a .docx")
    parser.add_argument("--docx", type=Path, required=True, help="Path to the .docx")
    parser.add_argument(
        "--figure-index", type=int, help="1-based figure index (1 = Figure 1, 2 = Figure 2, ...)"
    )
    parser.add_argument("--image", type=Path, help="Path to the replacement image")
    parser.add_argument("--list", action="store_true", help="List all figures + captions and exit")
    parser.add_argument(
        "--no-backup", action="store_true", help="Skip writing a timestamped backup of the .docx"
    )
    parser.add_argument(
        "--width-inches",
        type=float,
        default=None,
        help="Override target width (inches). Default fits to page content width.",
    )
    parser.add_argument(
        "--keep-size",
        action="store_true",
        help="Keep the original in-doc dimensions (do not resize).",
    )
    args = parser.parse_args()

    if not args.docx.is_file():
        parser.error(f"docx not found: {args.docx}")

    if args.list:
        list_figures(args.docx)
        return 0

    if args.figure_index is None or args.image is None:
        parser.error("--figure-index and --image are required unless --list is used")
    if not args.image.is_file():
        parser.error(f"image not found: {args.image}")

    replace_figure(
        docx_path=args.docx,
        figure_index=args.figure_index,
        new_image_path=args.image,
        backup=not args.no_backup,
        width_inches=args.width_inches,
        keep_size=args.keep_size,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
